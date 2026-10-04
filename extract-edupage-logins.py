#!/usr/bin/env python3
"""Pull the EduPage student logins out of an EduPage parent-letter PDF.

The PDF prints one letter per student: the student's name and class on one line,
then the student's username and password on the left with the parent e-mail
addresses on the right. `pdftotext -layout` keeps that two-column shape, which is
what this script reads.

    python3 extract-edupage-logins.py "Ellipse International School edupage3.pdf" \
        --grade 8 --xlsx Grade8_EduPage_Logins.xlsx --csv Grade8_EduPage_Logins.csv

With no --grade every class in the PDF is written out.
"""

import argparse
import csv
import re
import subprocess
import sys

# "   Abduraxmon Alimdjanov                    Grade 8A" — the class may read "6P-A"
NAME_RE = re.compile(r'^\s{2,}(\S.*?)\s{2,}Grade\s+(\d{1,2})([A-Za-z0-9-]*)\s*$')
# "   Username: AbduraxmonAlimdjanov           Father: hulkar@gmail.com"
USER_RE = re.compile(r'^\s*Username:\s*(\S+)(?:\s{2,}Father:\s*(.*?))?\s*$')
PASS_RE = re.compile(r'^\s*Password:\s*(\S+)(?:\s+\(password changed by the user\))?(?:\s{2,}Mother:\s*(.*?))?\s*$')
# EduPage blanks out a password the student has since changed.
CHANGED = '(changed by the user)'

FIELDS = ['Name', 'Class', 'Grade', 'Username', 'Password', 'Father', 'Mother']


def pdf_text(path):
    """The PDF as layout-preserved text."""
    try:
        out = subprocess.run(['pdftotext', '-layout', path, '-'],
                             capture_output=True, check=True)
    except FileNotFoundError:
        sys.exit('pdftotext not found — install poppler-utils')
    except subprocess.CalledProcessError as exc:
        sys.exit(f'pdftotext failed on {path}: {exc.stderr.decode()[:200]}')
    return out.stdout.decode('utf-8', 'replace')


def clean(value):
    """EduPage prints a bare dash where it has no e-mail address."""
    value = (value or '').strip()
    return '' if value in {'-', '--'} else value


def parse(text):
    """Every student letter in the text, in the order the PDF prints them."""
    students, current = [], None
    for line in text.splitlines():
        name = NAME_RE.match(line)
        if name:
            # A letter with no credentials printed is still a student.
            if current:
                students.append(current)
            label, grade, stream = name.group(1).strip(), name.group(2), name.group(3)
            current = {'Name': label, 'Class': grade + stream, 'Grade': int(grade),
                       'Username': '', 'Password': '', 'Father': '', 'Mother': ''}
            continue
        if not current:
            continue
        user = USER_RE.match(line)
        if user:
            current['Username'] = user.group(1)
            current['Father'] = clean(user.group(2))
            continue
        pwd = PASS_RE.match(line)
        if pwd:
            password = pwd.group(1)
            current['Password'] = CHANGED if set(password) == {'*'} else password
            current['Mother'] = clean(pwd.group(2))
    if current:
        students.append(current)
    return students


def check(students, source):
    """Complain about anything a human should look at before using the list."""
    problems = []
    for s in students:
        missing = [f for f in ('Username', 'Password') if not s[f]]
        if missing:
            problems.append(f'{s["Name"]} ({s["Class"]}): no {" or ".join(missing)} in the PDF')
        elif s['Password'] == CHANGED:
            problems.append(f'{s["Name"]} ({s["Class"]}): password already changed by the student, '
                            f'so the PDF does not print it')
    seen = {}
    for s in students:
        seen.setdefault(s['Username'], []).append(s['Name'])
    for username, names in seen.items():
        if username and len(names) > 1:
            problems.append(f'username {username} is shared by {", ".join(names)}')
    if not students:
        problems.append(f'no student letters found in {source}')
    return problems


def write_csv(path, students):
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(students)


def write_xlsx(path, students, title):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title

    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='2F5597')
    ws.append(FIELDS)
    for cell in ws[1]:
        cell.font, cell.fill = head, fill
        cell.alignment = Alignment(horizontal='center')

    for s in students:
        ws.append([s[f] for f in FIELDS])

    # A password is a string even when it reads as a number.
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = '@'

    widths = [26, 8, 7, 26, 14, 32, 32]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(FIELDS))}{ws.max_row}'
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('pdf', nargs='+', help='EduPage parent-letter PDF(s)')
    ap.add_argument('--grade', type=int, help='keep only this grade (e.g. 8)')
    ap.add_argument('--csv', help='write the list to this CSV file')
    ap.add_argument('--xlsx', help='write the list to this Excel file')
    args = ap.parse_args()

    students = []
    for path in args.pdf:
        found = parse(pdf_text(path))
        for problem in check(found, path):
            print(f'warning: {problem}', file=sys.stderr)
        students.extend(found)

    if args.grade is not None:
        students = [s for s in students if s['Grade'] == args.grade]
        if not students:
            sys.exit(f'no grade {args.grade} students in {", ".join(args.pdf)}')

    by_class = {}
    for s in students:
        by_class[s['Class']] = by_class.get(s['Class'], 0) + 1
    print(f'{len(students)} students: ' +
          ' · '.join(f'{c} {n}' for c, n in sorted(by_class.items())), file=sys.stderr)

    title = f'Grade {args.grade}' if args.grade is not None else 'Students'
    if args.csv:
        write_csv(args.csv, students)
        print(f'wrote {args.csv}', file=sys.stderr)
    if args.xlsx:
        write_xlsx(args.xlsx, students, title)
        print(f'wrote {args.xlsx}', file=sys.stderr)
    if not (args.csv or args.xlsx):
        writer = csv.DictWriter(sys.stdout, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(students)


if __name__ == '__main__':
    main()
