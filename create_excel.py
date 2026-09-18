import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.Workbook()

# Data organized by grade
data = {
    "Grade 7": [
        {"#": 1, "Name": "Abduraxlov Abduvosit", "Class": "7-D", "Subject 1": "Matematika", "Subject 2": "English", "Subject 3": "Fizika", "Subject 4": "", "Subject 5": ""},
        {"#": 2, "Name": "Anarbek Mirzaalizov", "Class": "7-A", "Subject 1": "Physics & Astronomy", "Subject 2": "Maths", "Subject 3": "Biology", "Subject 4": "English", "Subject 5": "Chemistry"},
        {"#": 3, "Name": "Ilyosxonov", "Class": "7-A", "Subject 1": "Math", "Subject 2": "Fizika", "Subject 3": "Biologiya", "Subject 4": "", "Subject 5": ""},
        {"#": 4, "Name": "Uzarboydi Imon", "Class": "7-B", "Subject 1": "English", "Subject 2": "Mathematics", "Subject 3": "ICT", "Subject 4": "Biology", "Subject 5": ""},
        {"#": 5, "Name": "Xosin Muhammadov", "Class": "7-D", "Subject 1": "Matematika", "Subject 2": "Fizika", "Subject 3": "English", "Subject 4": "", "Subject 5": ""},
        {"#": 6, "Name": "Kulboeva Medina-Aisha", "Class": "7-D", "Subject 1": "Math", "Subject 2": "English", "Subject 3": "", "Subject 4": "", "Subject 5": ""},
        {"#": 7, "Name": "Dexkanova Bibixonim", "Class": "7-D", "Subject 1": "Matematika", "Subject 2": "English", "Subject 3": "", "Subject 4": "", "Subject 5": ""},
        {"#": 8, "Name": "Parviz", "Class": "7 (?)", "Subject 1": "English", "Subject 2": "Matematika", "Subject 3": "Informatika", "Subject 4": "", "Subject 5": ""},
        {"#": 9, "Name": "Sharipov A.", "Class": "7 (?)", "Subject 1": "IT", "Subject 2": "English", "Subject 3": "Biology", "Subject 4": "", "Subject 5": ""},
    ],
    "Grade 8": [
        {"#": 1, "Name": "Tojiddinova Hadicha", "Class": "8-B", "Subject 1": "Math", "Subject 2": "English", "Subject 3": "History", "Subject 4": "Geography", "Subject 5": "Physics"},
        {"#": 2, "Name": "Boborova Gavxor", "Class": "8-B", "Subject 1": "Biologiya", "Subject 2": "English", "Subject 3": "Ximiya (Chemistry)", "Subject 4": "Matematika", "Subject 5": "Fizika"},
        {"#": 3, "Name": "Ziyodulla", "Class": "8-G", "Subject 1": "Tarix (History)", "Subject 2": "Matematika", "Subject 3": "Geografiya", "Subject 4": "Adabiyot (Literature)", "Subject 5": "Ingliz tili (English)"},
        {"#": 4, "Name": "Abdujabborov Behruz", "Class": "8-P", "Subject 1": "English", "Subject 2": "Biology", "Subject 3": "Chemistry", "Subject 4": "Global Perspective", "Subject 5": "Physics"},
        {"#": 5, "Name": "Xasanova Latofat", "Class": "8-P", "Subject 1": "Math", "Subject 2": "English", "Subject 3": "Physics", "Subject 4": "Biology", "Subject 5": "Russian"},
        {"#": 6, "Name": "Abdurashidova Mavludaxon", "Class": "8-P", "Subject 1": "Math", "Subject 2": "English", "Subject 3": "Physics", "Subject 4": "Biology", "Subject 5": "Russian"},
        {"#": 7, "Name": "Rashidov Ayyubxon", "Class": "8-P", "Subject 1": "Math", "Subject 2": "Physics", "Subject 3": "Chemistry", "Subject 4": "English", "Subject 5": "Uzbek / Russian"},
        {"#": 8, "Name": "Muhammad Sanan", "Class": "8-P", "Subject 1": "English", "Subject 2": "Biology", "Subject 3": "Chemistry", "Subject 4": "Global Perspective", "Subject 5": "Physics"},
        {"#": 9, "Name": "J. Umar", "Class": "8-P", "Subject 1": "Matematika", "Subject 2": "Robota texnika (Robotics)", "Subject 3": "English", "Subject 4": "Fizika", "Subject 5": "IT"},
        {"#": 10, "Name": "Ismailov Yusufxon", "Class": "8-G", "Subject 1": "English", "Subject 2": "Math", "Subject 3": "Tarix (History)", "Subject 4": "Kimyo (Chemistry)", "Subject 5": "Fizik (Physics)"},
        {"#": 11, "Name": "Shuhratov Johongir", "Class": "8-A", "Subject 1": "Mathematics", "Subject 2": "Physics", "Subject 3": "English", "Subject 4": "", "Subject 5": ""},
        {"#": 12, "Name": "Baxtiyorov Imron", "Class": "8-P", "Subject 1": "Math", "Subject 2": "English", "Subject 3": "Biology", "Subject 4": "Global Perspective", "Subject 5": "Uzbek language"},
        {"#": 13, "Name": "Muhammadali Kasimov", "Class": "8-G", "Subject 1": "English", "Subject 2": "Chemistry", "Subject 3": "Math", "Subject 4": "Geografiya", "Subject 5": ""},
    ],
    "Grade 9": [
        {"#": 1, "Name": "Isroilov Otabek", "Class": "9", "Subject 1": "Math", "Subject 2": "Tarix (History)", "Subject 3": "IT", "Subject 4": "Physics", "Subject 5": "English"},
        {"#": 2, "Name": "Anvar Maksumov", "Class": "9-B", "Subject 1": "Matematika", "Subject 2": "English", "Subject 3": "Russian", "Subject 4": "", "Subject 5": ""},
        {"#": 3, "Name": "Minkobilova Rayyona", "Class": "9-A", "Subject 1": "English", "Subject 2": "Geography", "Subject 3": "History", "Subject 4": "Biology", "Subject 5": "Chemistry"},
        {"#": 4, "Name": "Muxtarova Nurayyona", "Class": "9-A", "Subject 1": "English", "Subject 2": "Geography", "Subject 3": "History", "Subject 4": "Math", "Subject 5": ""},
        {"#": 5, "Name": "Oxtambaeva Noima", "Class": "9-G", "Subject 1": "Huquq (Law)", "Subject 2": "Fizika", "Subject 3": "History", "Subject 4": "Geography", "Subject 5": ""},
    ],
    "Grade 10": [
        {"#": 1, "Name": "Babosortaev Abdulla", "Class": "10-A", "Subject 1": "Fizika", "Subject 2": "Kimiya (Chemistry)", "Subject 3": "English", "Subject 4": "", "Subject 5": ""},
        {"#": 2, "Name": "Islomov Tongerin", "Class": "10-A", "Subject 1": "Huquq (Law)", "Subject 2": "OZI", "Subject 3": "", "Subject 4": "", "Subject 5": ""},
        {"#": 3, "Name": "Gulimov Bahodir", "Class": "10-A", "Subject 1": "Matematika", "Subject 2": "Fizika", "Subject 3": "English", "Subject 4": "", "Subject 5": ""},
        {"#": 4, "Name": "Mirinov Ahmad", "Class": "10 (?)", "Subject 1": "English", "Subject 2": "Chemistry", "Subject 3": "Biology", "Subject 4": "ICT", "Subject 5": "History"},
        {"#": 5, "Name": "Turaev Sardor", "Class": "10-G", "Subject 1": "Matematika", "Subject 2": "Math", "Subject 3": "", "Subject 4": "", "Subject 5": ""},
        {"#": 6, "Name": "Xudoynazarova Nodira", "Class": "10-B", "Subject 1": "Informatika", "Subject 2": "Matematika", "Subject 3": "Fizika", "Subject 4": "Biologiya", "Subject 5": ""},
        {"#": 7, "Name": "Bonijonova Bilboldolu", "Class": "10 (?)", "Subject 1": "English", "Subject 2": "Matematika", "Subject 3": "Biologiya", "Subject 4": "IT", "Subject 5": ""},
        {"#": 8, "Name": "Abdulboriy", "Class": "10-A", "Subject 1": "ICT", "Subject 2": "Biology", "Subject 3": "Geography", "Subject 4": "English", "Subject 5": "Chemistry"},
    ],
    "Grade 11": [
        {"#": 1, "Name": "Shukurullaeva Madinobonu", "Class": "11-A", "Subject 1": "English", "Subject 2": "Matematika", "Subject 3": "P.E.", "Subject 4": "Chemistry", "Subject 5": "Physics"},
        {"#": 2, "Name": "Kosimova Shofura", "Class": "11-G", "Subject 1": "Matematika", "Subject 2": "Ona tili (Native lang.)", "Subject 3": "P.E.", "Subject 4": "", "Subject 5": ""},
        {"#": 3, "Name": "Abdunazarova Xosiyat", "Class": "11-A", "Subject 1": "English", "Subject 2": "Tarix (History)", "Subject 3": "Matematika", "Subject 4": "Fizika", "Subject 5": "P.E."},
        {"#": 4, "Name": "Soliqbonu", "Class": "11-B", "Subject 1": "Matematika", "Subject 2": "Ona tili (Uzbek)", "Subject 3": "Ingliz tili (English)", "Subject 4": "Tarix (History)", "Subject 5": "ICT"},
        {"#": 5, "Name": "Anvarxonov Sardaloxon", "Class": "11-G", "Subject 1": "English", "Subject 2": "", "Subject 3": "", "Subject 4": "", "Subject 5": ""},
        {"#": 6, "Name": "Sherzixon Buxorov", "Class": "11-G", "Subject 1": "English", "Subject 2": "Fizika", "Subject 3": "Matematika", "Subject 4": "Kimyo (Chemistry)", "Subject 5": "Ingliz tili (English)"},
        {"#": 7, "Name": "Buzodirov Usmon", "Class": "11-A", "Subject 1": "Mathematics", "Subject 2": "Physics", "Subject 3": "English", "Subject 4": "Chemistry", "Subject 5": "Russian"},
    ],
}

# Styling
header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
title_font = Font(name="Calibri", bold=True, size=14, color="2E75B6")
data_font = Font(name="Calibri", size=11)
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

headers = ["#", "Student Name", "Class", "Subject 1", "Subject 2", "Subject 3", "Subject 4", "Subject 5"]

first_sheet = True
for grade_name, students in data.items():
    if first_sheet:
        ws = wb.active
        ws.title = grade_name
        first_sheet = False
    else:
        ws = wb.create_sheet(title=grade_name)

    # Title row
    ws.merge_cells("A1:H1")
    cell = ws["A1"]
    cell.value = f"Ellipse International School - Olympiad Students ({grade_name})"
    cell.font = title_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Data rows
    for row_idx, student in enumerate(students, 4):
        values = [
            student["#"],
            student["Name"],
            student["Class"],
            student["Subject 1"],
            student["Subject 2"],
            student["Subject 3"],
            student["Subject 4"],
            student["Subject 5"],
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 3):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Column widths
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 10
    for col_letter in ["D", "E", "F", "G", "H"]:
        ws.column_dimensions[col_letter].width = 22

# Summary sheet
ws_summary = wb.create_sheet(title="Summary", index=0)
ws_summary.merge_cells("A1:C1")
cell = ws_summary["A1"]
cell.value = "Ellipse International School - Olympiad Summary"
cell.font = title_font
cell.alignment = Alignment(horizontal="center")
ws_summary.row_dimensions[1].height = 30

summary_headers = ["Grade", "Number of Students", "Note"]
for col_idx, header in enumerate(summary_headers, 1):
    cell = ws_summary.cell(row=3, column=col_idx)
    cell.value = header
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

summary_data = [
    ("Grade 7", 9, "Some students missing grade/class info"),
    ("Grade 8", 13, ""),
    ("Grade 9", 5, ""),
    ("Grade 10", 8, "Some students missing grade/class info"),
    ("Grade 11", 7, ""),
    ("TOTAL", 42, ""),
]

for row_idx, (grade, count, note) in enumerate(summary_data, 4):
    for col_idx, val in enumerate([grade, count, note], 1):
        cell = ws_summary.cell(row=row_idx, column=col_idx)
        cell.value = val
        cell.font = data_font if grade != "TOTAL" else Font(name="Calibri", bold=True, size=11)
        cell.border = thin_border
        cell.alignment = center_align if col_idx != 3 else left_align

ws_summary.column_dimensions["A"].width = 15
ws_summary.column_dimensions["B"].width = 22
ws_summary.column_dimensions["C"].width = 40

output_path = "/home/user/test1/Ellipse_Olympiad_Students.xlsx"
wb.save(output_path)
print(f"Excel file saved to: {output_path}")
